import csv
import random
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
NLP_INTERIM_DIR = DATA_DIR / "nlp" / "interim"
NLP_FINAL_DIR = DATA_DIR / "nlp" / "final"

INPUT_XLSX = NLP_INTERIM_DIR / "labeling_sample_300_prelabeled.xlsx"
OUTPUT_ALL = NLP_INTERIM_DIR / "labeled_sample_all.csv"
OUTPUT_RELEVANT = NLP_FINAL_DIR / "labeled_sample_relevant.csv"
OUTPUT_TRAIN = NLP_FINAL_DIR / "train_labeled.csv"
OUTPUT_TEST = NLP_FINAL_DIR / "test_labeled.csv"
TEST_RATIO = 0.2
SEED = 42

GUIDE_TO_LABEL = {
    "label_guide_relevance": "relevance",
    "label_guide_sentiment": "sentiment",
    "label_guide_issue_type": "issue_type",
}

VALID_RELEVANCE = {"relevant", "not_relevant"}
VALID_SENTIMENT = {"positive", "neutral", "negative"}
VALID_ISSUE_TYPE = {
    "delay",
    "breakdown",
    "crowding",
    "aircond",
    "app_issue",
    "cleanliness",
    "staff_service",
    "fare_ticket",
    "information",
    "other",
}

RELEVANCE_FIXES = {
    "relvant": "relevant",
}


def normalize_value(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.upper() == "NA":
        return ""
    return text


def normalize_relevance(value: str) -> str:
    value = RELEVANCE_FIXES.get(value, value)
    return value if value in VALID_RELEVANCE else ""


def normalize_sentiment(value: str) -> str:
    return value if value in VALID_SENTIMENT else ""


def normalize_issue_type(value: str) -> str:
    return value if value in VALID_ISSUE_TYPE else ""


def read_rows_from_workbook(path: Path) -> list[dict]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active

    headers = [cell.value for cell in sheet[1]]
    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        rows.append({headers[i]: row[i] for i in range(len(headers))})
    return rows


def finalize_rows(rows: list[dict]) -> list[dict]:
    finalized = []
    for row in rows:
        clean_row = {k: normalize_value(v) for k, v in row.items()}

        clean_row["relevance"] = normalize_relevance(clean_row.get("label_guide_relevance", ""))
        clean_row["sentiment"] = normalize_sentiment(clean_row.get("label_guide_sentiment", ""))
        clean_row["issue_type"] = normalize_issue_type(clean_row.get("label_guide_issue_type", ""))

        if clean_row["relevance"] != "relevant":
            clean_row["sentiment"] = ""
            clean_row["issue_type"] = ""
            clean_row["route_line"] = ""

        finalized.append(clean_row)

    return finalized


def write_csv(path: Path, rows: list[dict]) -> None:
    if not rows:
        raise ValueError(f"No rows to write for {path.name}")

    fieldnames = list(rows[0].keys())
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def split_rows(rows: list[dict], test_ratio: float, seed: int) -> tuple[list[dict], list[dict]]:
    shuffled = rows[:]
    rng = random.Random(seed)
    rng.shuffle(shuffled)
    test_size = max(1, int(len(shuffled) * test_ratio))
    test_rows = shuffled[:test_size]
    train_rows = shuffled[test_size:]
    return train_rows, test_rows


def main() -> None:
    raw_rows = read_rows_from_workbook(INPUT_XLSX)
    finalized_rows = finalize_rows(raw_rows)
    relevant_rows = [row for row in finalized_rows if row.get("relevance") == "relevant"]
    train_rows, test_rows = split_rows(relevant_rows, TEST_RATIO, SEED)

    write_csv(OUTPUT_ALL, finalized_rows)
    write_csv(OUTPUT_RELEVANT, relevant_rows)
    write_csv(OUTPUT_TRAIN, train_rows)
    write_csv(OUTPUT_TEST, test_rows)

    print(f"Input workbook: {INPUT_XLSX.resolve()}")
    print(f"All labeled rows: {len(finalized_rows)} -> {OUTPUT_ALL.resolve()}")
    print(f"Relevant rows: {len(relevant_rows)} -> {OUTPUT_RELEVANT.resolve()}")
    print(f"Train rows: {len(train_rows)} -> {OUTPUT_TRAIN.resolve()}")
    print(f"Test rows: {len(test_rows)} -> {OUTPUT_TEST.resolve()}")


if __name__ == "__main__":
    main()
