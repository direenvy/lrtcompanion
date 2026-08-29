import csv
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
NLP_INTERIM_DIR = DATA_DIR / "nlp" / "interim"

INPUT_XLSX = NLP_INTERIM_DIR / "labeling_sample_300_prelabeled.xlsx"
OUTPUT_XLSX = NLP_INTERIM_DIR / "labeling_sample_300_human_checked.xlsx"
OUTPUT_CSV = NLP_INTERIM_DIR / "labeling_sample_300_human_checked.csv"

GUIDE_TO_LABEL = {
    "label_guide_relevance": "relevance",
    "label_guide_sentiment": "sentiment",
    "label_guide_issue_type": "issue_type",
}


def normalize_value(value: str) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.upper() == "NA":
        return ""
    return text


def main() -> None:
    workbook = load_workbook(INPUT_XLSX)
    sheet = workbook.active

    headers = [cell.value for cell in sheet[1]]
    column_index = {header: idx + 1 for idx, header in enumerate(headers) if header}

    missing = [name for name in GUIDE_TO_LABEL if name not in column_index]
    missing += [name for name in GUIDE_TO_LABEL.values() if name not in column_index]
    if missing:
        raise ValueError(f"Missing expected columns: {sorted(set(missing))}")

    updated_rows = 0
    for row in range(2, sheet.max_row + 1):
        row_changed = False
        for guide_col, label_col in GUIDE_TO_LABEL.items():
            guide_value = normalize_value(sheet.cell(row=row, column=column_index[guide_col]).value)
            if guide_value:
                sheet.cell(row=row, column=column_index[label_col]).value = guide_value
                row_changed = True

        if row_changed and "notes" in column_index:
            sheet.cell(row=row, column=column_index["notes"]).value = "human_checked"
            updated_rows += 1

    workbook.save(OUTPUT_XLSX)

    rows = list(sheet.values)
    with OUTPUT_CSV.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerows(rows)

    print(f"Input workbook: {INPUT_XLSX.resolve()}")
    print(f"Output workbook: {OUTPUT_XLSX.resolve()}")
    print(f"Output CSV: {OUTPUT_CSV.resolve()}")
    print(f"Rows updated: {updated_rows}")


if __name__ == "__main__":
    main()
